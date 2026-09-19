"""Build the Chinese introduction from current simulation results."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
pickup=json.loads((ROOT/'outputs/python_sim/pickup/result.json').read_text())
detour=json.loads((ROOT/'outputs/python_sim/detour/result.json').read_text())
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(21)
doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(12)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(6)
header=s.header.paragraphs[0];header.text='仓库机器人路径规划  |  Python 仿真介绍'
header.runs[0].font.size=Pt(9)
footer=s.footer.paragraphs[0];footer.alignment=2
footer.add_run('第 ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' 页')

def p(text):return doc.add_paragraph(text)
def h(text):return doc.add_heading(text,1)
def code(text):
    for line in text.splitlines():
        para=doc.add_paragraph();para.paragraph_format.space_after=Pt(3)
        run=para.add_run(line);run.font.name='Consolas';run.font.size=Pt(10)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    for row in rows:
        for cell,x in zip(t.add_row().cells,row):cell.text=str(x)
    for ri,row in enumerate(t.rows):
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=1
            tcPr=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E4EBF2' if ri==0 else ('F5F7FA' if ri%2==0 else 'FFFFFF'));tcPr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcPr.append(borders)
            for para in cell.paragraphs:
                para.paragraph_format.space_before=Pt(5);para.paragraph_format.space_after=Pt(5)
                for r in para.runs:r.font.size=Pt(10);r.bold=ri==0
    # Table ends with its own paragraph; no empty spacer paragraph is needed.

doc.add_heading('纯 Python 仓库机器人\nA 星路径规划仿真介绍',0)
p('项目说明与演示指南    2026 年 9 月 19 日')
p('本项目在 Python 中复现仓库地图和机器人通行约束，用 A* 算法搜索路径，并以动画展示搜索过程与机器人沿路径运动。它适合算法学习、仓库布局验证和项目演示，无需启动 Unity、ROS 或 Docker。')
h('项目解决的问题')
p('仓库小车能否通过货架下方，不仅取决于平面位置，还取决于车身朝向、举升高度和是否载货。因此，仿真同时考虑低位交互通道、高净空通道、货架立柱及临时障碍，避免把机器人简化成没有尺寸的点。')
p('目前提供两个可直接演示的场景：空载机器人从起点驶向 P1 交互平台，以及载货机器人在高通道遇到障碍后绕行。动画用于解释路径如何得到；每条路径还会单独检查几何碰撞。')
h('地图与机器人定义')
table(['项目','当前设置'],[
('地图范围','18 m × 20 m，栅格分辨率 0.025 m'),
('坐标约定','map 坐标，距离单位为米，航向角为弧度'),
('搜索状态','栅格列、栅格行和 8 个离散朝向'),
('机器人轮廓','底盘及车轮约 0.90 m × 0.956 m'),
('载货轮廓','负载平面尺寸 1.70 m × 1.40 m'),
('机器人模式','空载、升高空载、载货；当前动画展示空载与载货'),
('规划余量','平面安全余量 0.02 m')],[1.5,5.4])
p('地图来源是已导出的仓库几何快照。独立运行仅需规划代码、配置和地图文件，不需要加载 Unity 工程。')

doc.add_page_break();h('A 星算法如何规划路径')
p('A* 在候选状态中优先扩展估计总代价较小的状态。g 表示从起点到当前状态已经付出的代价，h 表示到目标的剩余代价估计，两者之和 f 用于决定搜索顺序。')
doc.add_heading('动作与代价',2)
p('机器人可以沿当前朝向前进一格、倒退一格，或者原地转向 45°。算法不允许横向平移；斜向移动会检查相邻栅格，防止从障碍物角落穿过。转向需要通过旋转扫掠检查，P1 和 P2 的交互槽内额外禁止原地转向。')
p('前进代价等于移动距离；倒车距离乘以 1.1；每次 45° 转向增加 0.18 的代价。启发函数采用适合八邻域栅格的 octile 距离。最终求得的是当前离散动作图中的最小代价路径，不代表连续空间中任意曲线的全局最优。')
doc.add_heading('搜索到动画的流程',2)
p('读取几何快照和配置 → 选择机器人模式 → 加载或生成朝向相关的障碍地图 → 执行 A* 并记录展开顺序 → 回溯最终路径 → 独立碰撞校验 → 生成动画及结果文件。')
p('搜索画面取自真实的 A* 展开回调，而不是预制的扩散特效。为便于观察，同一 XY 栅格只显示第一次展开，并分批压缩播放；算法内部仍保留各个朝向状态。画面没有单独显示开放列表。')
h('验证结果')
table(['场景','路径长度','展开状态数','几何检查'],[
('空载起点到 P1',f'{pickup["length_m"]:.2f} m',f'{pickup["expansions"]:,}','通过'),
('载货高通道绕障',f'{detour["length_m"]:.2f} m',f'{detour["expansions"]:,}','通过')],[2.5,1.3,1.7,1.4])
p('当前记录包含 18 项通过的单元测试，覆盖 A* 与 Dijkstra 代价一致性、倒车与转向、障碍与高度约束、路径跟踪基础逻辑，以及搜索记录不改变路径和动画转向插值。')
p('独立碰撞检查按最多 0.0125 m 的平移步长和 2° 的转向步长抽查机器人各高度层。动画自身采用更稀疏的帧间插值来提高可读性，播放速度不等于真实机器人速度。')

doc.add_page_break();h('动画示例与阅读方法')
p('下图展示载货机器人遇到临时障碍后的绕行结果。实际播放器先展示青色的搜索展开区域，再显示绿色路径和橙色机器人沿路径运动。')
para=doc.add_paragraph();para.alignment=1
para.add_run().add_picture(str(ROOT/'outputs/python_sim/detour/preview.png'),width=Inches(4.9))
p('图 1  载货绕障场景的最终画面  动画文件位于 outputs/python_sim/detour')
p('深灰表示低处实体，金色表示高处结构，红色表示临时障碍。金色区域并非一律禁止通过：是否碰撞由机器人高度、负载和朝向共同决定。橙色轮廓及箭头表达机器人占地范围与前进方向。')
p('播放器支持暂停、重播、进度拖动与速度切换。Word 中的图片是静态示例；请打开输出目录里的 player.html 观看交互动画，或打开 simulation.gif 循环播放。')

doc.add_page_break();h('直接使用 Python 运行')
p('以下命令输入在 PowerShell 或 VS Code 的终端中，不是在 Python 的 >>> 提示符中。先进入项目根目录，再运行 Python 模块，无需执行 CMD 或 PS1 启动脚本。')
code('cd "C:\\Users\\hairo\\OneDrive\\文档\\找工\\load-aware-warehouse-robot"')
doc.add_heading('使用本机已验证的解释器',2)
code('$simPython = "$env:USERPROFILE\\.cache\\codex-runtimes\\codex-primary-runtime\\dependencies\\python\\python.exe"\n& $simPython --version\n& $simPython -c "import numpy, PIL; print(\'Dependencies OK\')"')
p('如果使用自行安装且可通过 python 命令访问的解释器，首次安装依赖，然后执行以下命令。依赖已安装时无需重复安装。')
code('python -m pip install -r requirements-planning.txt\npython -m warehouse_planning.animation --scenario pickup --show\npython -m warehouse_planning.animation --scenario detour --show')
p('使用上面定义的本机解释器时，把命令开头的 python 换为 & $simPython。例如：')
code('& $simPython -m warehouse_planning.animation --scenario pickup --show')
p('程序先完成搜索和动画生成，再自动打开默认浏览器。去掉 --show 只生成文件。终端显示 geometry PASS 代表本次路径通过几何检查；生成结束后可关闭终端，浏览器播放器仍可独立播放。')
doc.add_heading('输出文件',2)
table(['文件','用途'],[('player.html','离线交互播放器，无需网络'),('simulation.gif','可分享的循环动画'),('preview.png','最终路径静态图'),('result.json','路径、长度、展开数量和验证结果')],[2,4.9])
p('输出目录为 outputs/python_sim/pickup 或 outputs/python_sim/detour。完整运行步骤与 Python 或 Jupyter 调用示例见 docs/PYTHON_SIMULATION.md。')
doc.add_heading('仿真范围与后续扩展',2)
p('本版本是几何规划与运动回放，不模拟轮地摩擦、惯性和传感器噪声，不执行货物连接与释放。ROS–Unity 物理闭环是项目中的另一条验证流程，不能用本动画结果替代其执行报告。后续可扩展交互式目标点、障碍物编辑和不同代价策略对比。')
for element in [doc.styles.element,doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for paragraph in doc.paragraphs:
    prop=paragraph._p.get_or_add_pPr()
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');prop.append(snap)
doc.save(ROOT/'docs/纯Python仓库A星路径规划仿真介绍.docx')
print('DOCX created')
